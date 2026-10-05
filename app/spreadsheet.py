"""Read employees from, and write results back to, the Excel workbook.

Workbook layout (see sample_data/7X_Employees_SAMPLE.xlsx):
  * Sheet "Employees" - one row per employee, header in row 1.
  * Sheet "Event_Log" - one row per message attempt (mirror of the SQLite log).

Columns are located by header name, so their order may change.
Writes are atomic: the workbook is saved to a temp file and then renamed.
"""

from __future__ import annotations

import os
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.models import (
    EMPLOYEE_COLUMNS,
    EVENT_LOG_COLUMNS,
    Employee,
    normalize_phone,
    parse_date,
    parse_year,
    parse_yes,
)

REQUIRED_COLUMNS = [c for c in EMPLOYEE_COLUMNS if c != "Language"]


class SpreadsheetUnavailable(RuntimeError):
    """The workbook cannot be opened, read or saved."""


class EmployeeWorkbook:
    def __init__(self, path: Path, employees_sheet: str = "Employees",
                 event_log_sheet: str = "Event_Log"):
        self.path = Path(path)
        self.employees_sheet = employees_sheet
        self.event_log_sheet = event_log_sheet
        if not self.path.exists():
            raise SpreadsheetUnavailable(f"Spreadsheet not found: {self.path}")
        try:
            self.wb = load_workbook(self.path)
        except Exception as exc:  # corrupt file, locked file, wrong format ...
            raise SpreadsheetUnavailable(f"Cannot open spreadsheet {self.path.name}: {exc}") from exc
        if employees_sheet not in self.wb.sheetnames:
            raise SpreadsheetUnavailable(f"Sheet '{employees_sheet}' not found in {self.path.name}")
        self.ws = self.wb[employees_sheet]
        self.columns = self._read_header()
        missing = [c for c in REQUIRED_COLUMNS if c.lower() not in self.columns]
        if missing:
            raise SpreadsheetUnavailable(f"Missing required column(s): {', '.join(missing)}")

    # ----------------------------------------------------------------- header
    def _read_header(self) -> dict[str, int]:
        cols = {}
        for idx, cell in enumerate(self.ws[1], start=1):
            if cell.value is not None and str(cell.value).strip():
                cols[str(cell.value).strip().lower()] = idx
        return cols

    def _get(self, row: int, column: str):
        idx = self.columns.get(column.lower())
        return self.ws.cell(row=row, column=idx).value if idx else None

    def _set(self, row: int, column: str, value) -> None:
        idx = self.columns.get(column.lower())
        if idx:
            self.ws.cell(row=row, column=idx).value = value

    # ------------------------------------------------------------------ reads
    def read_employees(self) -> list[Employee]:
        employees: list[Employee] = []
        for row in range(2, self.ws.max_row + 1):
            emp_id = self._get(row, "Employee_ID")
            name = self._get(row, "Name")
            if (emp_id is None or str(emp_id).strip() == "") and not name:
                continue  # blank row
            employees.append(self._parse_row(row, emp_id, name))
        return employees

    def _parse_row(self, row: int, emp_id, name) -> Employee:
        phone_raw = self._get(row, "Phone")
        ok, phone, phone_err = normalize_phone(phone_raw)
        birthday, b_err = parse_date(self._get(row, "Birthday"))
        joining, j_err = parse_date(self._get(row, "Joining_Date"))
        language = str(self._get(row, "Language") or "EN").strip().upper() or "EN"
        return Employee(
            row_number=row,
            employee_id=str(emp_id or f"ROW{row}").strip(),
            name=str(name or "").strip(),
            phone_raw="" if phone_raw is None else str(phone_raw),
            phone=phone if ok else None,
            phone_error=phone_err,
            birthday=birthday,
            birthday_error=None if birthday else f"Birthday {b_err}",
            joining_date=joining,
            joining_error=None if joining else f"Joining_Date {j_err}",
            consent=parse_yes(self._get(row, "WhatsApp_Consent")),
            active=parse_yes(self._get(row, "Active")),
            language=language if language in ("EN", "AR") else "EN",
            birthday_sent_year=parse_year(self._get(row, "Birthday_Sent_Year")),
            anniversary_sent_year=parse_year(self._get(row, "Anniversary_Sent_Year")),
        )

    # ----------------------------------------------------------------- writes
    def record_result(self, employee: Employee, *, message_type: str | None = None,
                      status: str, error: str = "", sent_year_column: str | None = None,
                      year: int | None = None, when: datetime | None = None) -> None:
        row = employee.row_number
        if sent_year_column and year is not None:
            self._set(row, sent_year_column, year)
        if message_type:
            self._set(row, "Last_Message_Type", message_type)
            self._set(row, "Last_Message_Date", (when or datetime.now()).strftime("%Y-%m-%d %H:%M"))
        self._set(row, "Send_Status", status)
        self._set(row, "Error_Message", (error or "")[:300])

    def append_event_log(self, values: dict) -> None:
        if self.event_log_sheet not in self.wb.sheetnames:
            ws = self.wb.create_sheet(self.event_log_sheet)
            ws.append(EVENT_LOG_COLUMNS)
        ws = self.wb[self.event_log_sheet]
        ws.append([values.get(col, "") for col in EVENT_LOG_COLUMNS])

    def save(self) -> None:
        """Atomic save: write a temp file next to the workbook, then replace."""
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=self.path.parent)
        os.close(fd)
        try:
            self.wb.save(tmp)
            os.replace(tmp, self.path)
        except Exception as exc:
            if os.path.exists(tmp):
                os.remove(tmp)
            raise SpreadsheetUnavailable(
                f"Cannot save spreadsheet {self.path.name} (is it open in Excel?): {exc}"
            ) from exc
