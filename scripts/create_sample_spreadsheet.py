"""Create the employee workbook template / sample (FAKE employees only).

    python scripts/create_sample_spreadsheet.py                 # sample with fake people
    python scripts/create_sample_spreadsheet.py --empty OUT.xlsx  # clean template, no rows

The workbook contains:
  * Employees  - data + validation rules + helper formula columns for Make.com
  * Event_Log  - audit / idempotency log
  * Lists      - allowed values for drop-downs
  * README     - how to fill in the sheet
"""

from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.models import EMPLOYEE_COLUMNS, EVENT_LOG_COLUMNS, HELPER_COLUMNS  # noqa: E402

MAX_ROWS = 500  # validation rules are applied to rows 2..MAX_ROWS
HELPER_ROWS = 200  # helper formulas pre-filled for this many data rows

# Fake people only. Phone numbers use the reserved-looking +97150000xxxx range.
FAKE_EMPLOYEES = [
    ("EMP001", "Ahmed Example", "+971500000001", date(1980, 3, 14), date(2015, 6, 1), "YES", "YES", "EN"),
    ("EMP002", "Sara Sample", "+971500000002", date(1985, 7, 22), date(2022, 9, 18), "YES", "YES", "EN"),
    ("EMP003", "Omar Placeholder", "+971500000003", date(1978, 11, 2), date(2010, 1, 10), "YES", "YES", "AR"),
    ("EMP004", "Layla Demo", "+971500000004", date(1992, 2, 29), date(2020, 2, 29), "YES", "YES", "EN"),
    ("EMP005", "Khalid Fictional", "+971500000005", date(1975, 5, 5), date(2008, 5, 5), "NO", "YES", "EN"),
    ("EMP006", "Mona Testcase", "+971500000006", date(1988, 12, 25), date(2018, 4, 15), "YES", "NO", "EN"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HELPER_FILL = PatternFill("solid", fgColor="7F7F7F")
SYSTEM_FILL = PatternFill("solid", fgColor="375623")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
SYSTEM_COLUMNS = {"Birthday_Sent_Year", "Anniversary_Sent_Year", "Last_Message_Type",
                  "Last_Message_Date", "Send_Status", "Error_Message"}


def col(name: str, headers: list[str]) -> str:
    return get_column_letter(headers.index(name) + 1)


def helper_formulas(r: int, headers: list[str]) -> dict[str, str]:
    b, j = col("Birthday", headers), col("Joining_Date", headers)
    # 29 Feb is celebrated on 28 Feb in non-leap years (same rule as the Python code).
    leap_fix = 'IF(AND(MONTH({c}{r})=2,DAY({c}{r})=29,DAY(DATE(YEAR(TODAY()),3,0))=28),"02-28",' \
               'TEXT(MONTH({c}{r}),"00")&"-"&TEXT(DAY({c}{r}),"00"))'
    return {
        "Birthday_MMDD": f'=IF({b}{r}="","",' + leap_fix.format(c=b, r=r) + ")",
        "Joining_MMDD": f'=IF({j}{r}="","",' + leap_fix.format(c=j, r=r) + ")",
        "Joining_Year": f'=IF({j}{r}="","",YEAR({j}{r}))',
        "Row_No": "=ROW()",
    }


def build(path: Path, with_samples: bool) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    headers = EMPLOYEE_COLUMNS + HELPER_COLUMNS
    ws.append(headers)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = WHITE_BOLD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.fill = HELPER_FILL if h in HELPER_COLUMNS else SYSTEM_FILL if h in SYSTEM_COLUMNS else HEADER_FILL
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 4)
    ws.column_dimensions[col("Name", headers)].width = 24
    ws.column_dimensions[col("Error_Message", headers)].width = 40
    ws.freeze_panes = "C2"

    rows = FAKE_EMPLOYEES if with_samples else []
    for r, (eid, name, phone, bday, joined, consent, active, lang) in enumerate(rows, start=2):
        values = {"Employee_ID": eid, "Name": name, "Phone": phone, "Birthday": bday,
                  "Joining_Date": joined, "WhatsApp_Consent": consent, "Active": active,
                  "Language": lang}
        for h, v in values.items():
            ws.cell(row=r, column=headers.index(h) + 1, value=v)

    # Formats for every data row (and helper formulas pre-filled)
    for r in range(2, MAX_ROWS + 1):
        ws[f"{col('Phone', headers)}{r}"].number_format = "@"  # keep the leading +
        ws[f"{col('Birthday', headers)}{r}"].number_format = "yyyy-mm-dd"
        ws[f"{col('Joining_Date', headers)}{r}"].number_format = "yyyy-mm-dd"
        if r <= HELPER_ROWS + 1:
            for h, f in helper_formulas(r, headers).items():
                ws[f"{col(h, headers)}{r}"] = f

    # ---------------- validation rules
    rng = lambda name: f"{col(name, headers)}2:{col(name, headers)}{MAX_ROWS}"  # noqa: E731
    p = col("Phone", headers)
    e = col("Employee_ID", headers)
    rules = [
        DataValidation(type="list", formula1='"YES,NO"', allow_blank=False,
                       errorTitle="YES or NO", error="Type YES or NO.", sqref=rng("WhatsApp_Consent")),
        DataValidation(type="list", formula1='"YES,NO"', allow_blank=False,
                       errorTitle="YES or NO", error="Type YES or NO.", sqref=rng("Active")),
        DataValidation(type="list", formula1='"EN,AR"', allow_blank=True,
                       errorTitle="Language", error="Type EN or AR (blank = EN).", sqref=rng("Language")),
        DataValidation(type="date", operator="between", formula1="DATE(1930,1,1)",
                       formula2="TODAY()", errorTitle="Birthday",
                       error="Enter a real date (e.g. 1985-07-22).", sqref=rng("Birthday")),
        DataValidation(type="date", operator="between", formula1="DATE(1970,1,1)",
                       formula2="TODAY()", errorTitle="Joining date",
                       error="Enter a real date that is not in the future.", sqref=rng("Joining_Date")),
        DataValidation(type="custom",
                       formula1=f'AND(LEFT({p}2,1)="+",LEN({p}2)>=9,LEN({p}2)<=16,'
                                f'ISNUMBER(VALUE(MID({p}2,2,15))))',
                       errorTitle="Phone format",
                       error="Use international format with no spaces, e.g. +971501234567",
                       sqref=rng("Phone")),
        DataValidation(type="custom", formula1=f"COUNTIF(${e}$2:${e}${MAX_ROWS},{e}2)=1",
                       errorTitle="Duplicate ID", error="Employee_ID must be unique.",
                       sqref=rng("Employee_ID")),
    ]
    for dv in rules:
        dv.showErrorMessage = True
        ws.add_data_validation(dv)

    red = PatternFill("solid", fgColor="F8CBAD")
    ws.conditional_formatting.add(
        rng("Phone"),
        FormulaRule(formula=[f'AND({p}2<>"",OR(LEFT({p}2,1)<>"+",LEN({p}2)<9))'], fill=red))
    ws.conditional_formatting.add(
        rng("Send_Status"),
        FormulaRule(formula=[f'OR({col("Send_Status", headers)}2="FAILED",'
                             f'{col("Send_Status", headers)}2="UNKNOWN",'
                             f'{col("Send_Status", headers)}2="DATA_ERROR")'], fill=red))

    # ---------------- Event_Log
    log_ws = wb.create_sheet("Event_Log")
    log_ws.append(EVENT_LOG_COLUMNS)
    for i, h in enumerate(EVENT_LOG_COLUMNS, start=1):
        c = log_ws.cell(row=1, column=i)
        c.font, c.fill = WHITE_BOLD, HEADER_FILL
        log_ws.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 6)
    log_ws.column_dimensions["A"].width = 32
    log_ws.freeze_panes = "A2"

    # ---------------- Lists
    lists = wb.create_sheet("Lists")
    lists.append(["YES_NO", "LANGUAGE", "SEND_STATUS"])
    for row in (["YES", "EN", "SENT"], ["NO", "AR", "TEST_SENT"], ["", "", "FAILED"],
                ["", "", "UNKNOWN"], ["", "", "DATA_ERROR"]):
        lists.append(row)

    # ---------------- README sheet
    readme = wb.create_sheet("README")
    lines = [
        "7X Executive Connect - how to fill in the Employees sheet",
        "",
        "BLUE columns: you fill these in.",
        "GREEN columns: written by the automation - do not edit unless correcting a mistake.",
        "GREY columns: formulas used by Make.com - copy them down for new rows, never type in them.",
        "",
        "Employee_ID: unique code, e.g. EMP001. Never reuse an ID.",
        "Phone: international format, no spaces: +971501234567",
        "Birthday / Joining_Date: real Excel dates (e.g. 1985-07-22). Birth year is ignored.",
        "WhatsApp_Consent: YES only if the employee agreed to receive WhatsApp greetings.",
        "Active: YES for current employees. Set NO when someone leaves (do not delete the row).",
        "Language: EN or AR (blank = EN). AR is used only after Arabic templates are approved.",
        "",
        "A message is sent ONLY when WhatsApp_Consent = YES AND Active = YES.",
        "Each person receives each celebration at most once per year.",
        "",
        "This workbook contains personal data. Store it only in the approved 7X OneDrive/SharePoint.",
    ]
    for line in lines:
        readme.append([line])
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 110

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--empty", metavar="OUT", help="create an EMPTY template at this path")
    ap.add_argument("--out", default=str(ROOT / "sample_data" / "7X_Employees_SAMPLE.xlsx"))
    args = ap.parse_args()
    if args.empty:
        print("Created", build(Path(args.empty), with_samples=False))
    else:
        print("Created", build(Path(args.out), with_samples=True))


if __name__ == "__main__":
    main()
