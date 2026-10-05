from datetime import date

from openpyxl import load_workbook

from app.birthday import birthday_event, is_birthday_today
from app.main import run_daily
from app.models import Employee
from conftest import TEST_PHONE, TODAY, employee_row, make_workbook


def emp(**kw) -> Employee:
    base = dict(row_number=2, employee_id="EMP001", name="Ahmed Example", phone="+971500000001",
                consent=True, active=True)
    base.update(kw)
    return Employee(**base)


# 1. Birthday today
def test_birthday_today_is_detected():
    e = emp(birthday=date(1980, 10, 5))
    ev = birthday_event(e, TODAY)
    assert ev is not None
    assert ev.event_id == "EMP001_BIRTHDAY_2026"


def test_birth_year_is_ignored():
    assert is_birthday_today(date(1999, 10, 5), TODAY)
    assert is_birthday_today(date(1950, 10, 5), TODAY)


# 2. Birthday tomorrow
def test_birthday_tomorrow_not_sent():
    assert birthday_event(emp(birthday=date(1980, 10, 6)), TODAY) is None


# 3. Already sent this year
def test_birthday_already_sent_this_year():
    assert birthday_event(emp(birthday=date(1980, 10, 5), birthday_sent_year=2026), TODAY) is None


def test_birthday_sent_last_year_is_sent_again():
    assert birthday_event(emp(birthday=date(1980, 10, 5), birthday_sent_year=2025), TODAY) is not None


def test_feb_29_birthday_celebrated_on_feb_28_in_non_leap_year():
    assert is_birthday_today(date(1992, 2, 29), date(2027, 2, 28))
    assert not is_birthday_today(date(1992, 2, 29), date(2027, 3, 1))
    assert is_birthday_today(date(1992, 2, 29), date(2028, 2, 29))
    assert not is_birthday_today(date(1992, 2, 29), date(2028, 2, 28))


def test_missing_birthday_reported_but_anniversary_still_works(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [
        employee_row(birthday=None, joining=date(2020, 10, 5)),
    ])
    s = run_daily(settings, TODAY, wa)
    assert any("Birthday missing" in d for d in s.data_errors)
    assert len(s.sent) == 1 and "Anniversary" in s.sent[0]


def test_invalid_birthday_text_reported(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday="31/02/1990")])
    s = run_daily(settings, TODAY, wa)
    assert any("invalid date" in d for d in s.data_errors)
    assert fake_session.calls == []


def test_birthday_end_to_end_production_updates_sheet(settings, wa, fake_session):
    from app.config import PRODUCTION_CONFIRM_PHRASE

    settings.app_mode = "PRODUCTION"
    settings.production_confirm = PRODUCTION_CONFIRM_PHRASE
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=date(1980, 10, 5))])
    s = run_daily(settings, TODAY, wa)
    assert len(s.sent) == 1
    assert fake_session.recipients == ["971500000001"]  # production -> real number
    ws = load_workbook(settings.spreadsheet_path)["Employees"]
    header = [c.value for c in ws[1]]
    row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(header)}
    assert row["Birthday_Sent_Year"] == 2026
    assert row["Last_Message_Type"] == "Birthday"
    assert row["Send_Status"] == "SENT"
    assert row["Last_Message_Date"]
    template = fake_session.calls[0]["json"]["template"]
    assert template["name"] == "employee_birthday"
    assert template["components"][0]["parameters"] == [{"type": "text", "text": "Test Person"}]
    assert TEST_PHONE.lstrip("+") not in fake_session.recipients
