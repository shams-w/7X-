"""Idempotency: the same celebration must never be sent twice in a year."""

from datetime import date

import requests
from openpyxl import load_workbook

from app.event_log import ClaimOutcome, EventLog
from app.main import run_daily
from app.models import CelebrationEvent, Employee, EventType
from app.whatsapp import WhatsAppService
from conftest import TODAY, FakeResponse, FakeSession, employee_row, make_workbook

BDAY = date(1980, 10, 5)


# 12. Duplicate automation run
def test_second_run_same_day_sends_nothing(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY)])
    first = run_daily(settings, TODAY, wa)
    second = run_daily(settings, TODAY, wa)
    assert len(first.sent) == 1
    assert second.sent == []
    assert len(second.duplicates) == 1
    assert len(fake_session.calls) == 1


def test_three_runs_with_new_service_instances(settings):
    """Simulates script restarts / manual reruns (fresh objects each time)."""
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY, joining=date(2020, 10, 5))])
    sessions = []
    for _ in range(3):
        sess = FakeSession()
        sessions.append(sess)
        run_daily(settings, TODAY, WhatsAppService(settings, session=sess, sleep=lambda s: None))
    assert [len(s.calls) for s in sessions] == [2, 0, 0]


def test_event_id_format():
    e = Employee(row_number=2, employee_id="emp001", name="X")
    assert CelebrationEvent(e, EventType.BIRTHDAY, 2026, TODAY).event_id == "EMP001_BIRTHDAY_2026"
    assert CelebrationEvent(e, EventType.ANNIVERSARY, 2026, TODAY).event_id == "EMP001_ANNIVERSARY_2026"


def test_claim_is_atomic(settings):
    log = EventLog(settings.event_db_path, settings.audit_csv_path, "TEST")
    ev = CelebrationEvent(Employee(row_number=2, employee_id="EMP001", name="X"),
                          EventType.BIRTHDAY, 2026, TODAY)
    assert log.claim(ev, "+971500000001") is ClaimOutcome.CLAIMED
    # A crashed/parallel run sees PENDING and must not send.
    assert log.claim(ev, "+971500000001") is ClaimOutcome.IN_PROGRESS_OR_UNKNOWN
    log.mark_sent(ev.event_id, "wamid.1")
    assert log.claim(ev, "+971500000001") is ClaimOutcome.ALREADY_SENT


def test_failed_send_can_be_retried_later(settings):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY)])
    failing = FakeSession([FakeResponse(400, {"error": {"code": 131000, "message": "Something went wrong"}})])
    s1 = run_daily(settings, TODAY, WhatsAppService(settings, session=failing, sleep=lambda s: None))
    assert len(s1.failed) == 1
    ok = FakeSession()
    s2 = run_daily(settings, TODAY, WhatsAppService(settings, session=ok, sleep=lambda s: None))
    assert len(s2.sent) == 1 and len(ok.calls) == 1


def test_timeout_after_send_is_never_auto_resent(settings):
    """Network timeout: Meta may have delivered the message -> mark UNKNOWN, do not resend."""
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY)])
    timeout = FakeSession([requests.ReadTimeout("read timed out")])
    s1 = run_daily(settings, TODAY, WhatsAppService(settings, session=timeout, sleep=lambda s: None))
    assert len(s1.unknown) == 1
    again = FakeSession()
    s2 = run_daily(settings, TODAY, WhatsAppService(settings, session=again, sleep=lambda s: None))
    assert again.calls == [] and len(s2.unknown) == 1
    # Operator confirms it was NOT delivered and resets it -> next run sends.
    assert EventLog(settings.event_db_path, settings.audit_csv_path, "TEST").reset("EMP001_BIRTHDAY_2026")
    s3 = run_daily(settings, TODAY, WhatsAppService(settings, session=again, sleep=lambda s: None))
    assert len(s3.sent) == 1


def test_connection_refused_is_retried_within_run(settings):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY)])
    sess = FakeSession([requests.ConnectionError("NewConnectionError: Failed to establish a new connection")])
    s = run_daily(settings, TODAY, WhatsAppService(settings, session=sess, sleep=lambda s: None))
    assert len(s.sent) == 1 and len(sess.calls) == 2


def test_spreadsheet_sent_year_also_blocks(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY, Birthday_Sent_Year=2026)])
    s = run_daily(settings, TODAY, wa)
    assert fake_session.calls == [] and len(s.duplicates) == 1


def test_test_mode_does_not_mark_real_sent_year(settings, wa):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY)])
    run_daily(settings, TODAY, wa)
    ws = load_workbook(settings.spreadsheet_path)["Employees"]
    header = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=header.index("Birthday_Sent_Year") + 1).value is None
    assert ws.cell(row=2, column=header.index("Send_Status") + 1).value == "TEST_SENT"
    log_ws = load_workbook(settings.spreadsheet_path)["Event_Log"]
    assert log_ws.cell(row=2, column=1).value == "EMP001_BIRTHDAY_2026"


def test_test_and_production_namespaces_are_separate(settings):
    ev = CelebrationEvent(Employee(row_number=2, employee_id="EMP001", name="X"),
                          EventType.BIRTHDAY, 2026, TODAY)
    test_log = EventLog(settings.event_db_path, settings.audit_csv_path, "TEST")
    test_log.claim(ev, "+1")
    test_log.mark_sent(ev.event_id, "wamid.t")
    prod_log = EventLog(settings.event_db_path, settings.audit_csv_path, "PRODUCTION")
    assert prod_log.check(ev.event_id) is ClaimOutcome.CLAIMED
