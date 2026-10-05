"""Eligibility, error handling, TEST mode, DRY RUN and secret-handling tests."""

from datetime import date

import pytest

from app.config import PRODUCTION_CONFIRM_PHRASE, ConfigError, Settings
from app.main import run_daily
from app.models import mask_phone, normalize_phone
from app.whatsapp import WhatsAppService
from conftest import TEST_PHONE, TODAY, FakeResponse, FakeSession, employee_row, make_workbook

BDAY = date(1980, 10, 5)


def svc(settings, session):
    return WhatsAppService(settings, session=session, sleep=lambda s: None)


# 9. No consent
def test_no_consent_not_sent(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY, consent="NO")])
    s = run_daily(settings, TODAY, wa)
    assert fake_session.calls == []
    assert any("no WhatsApp consent" in x for x in s.skipped)


def test_blank_consent_treated_as_no(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY, consent=None)])
    run_daily(settings, TODAY, wa)
    assert fake_session.calls == []


# 10. Inactive employee
def test_inactive_not_sent(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY, active="NO")])
    s = run_daily(settings, TODAY, wa)
    assert fake_session.calls == []
    assert any("not Active" in x for x in s.skipped)


# 11. Invalid phone
@pytest.mark.parametrize("phone", ["0501234567", "+97150", "abc", "", None, "+971 50 12A 4567"])
def test_invalid_phone_rejected(phone):
    ok, _, err = normalize_phone(phone)
    assert not ok and err


@pytest.mark.parametrize("raw,expected", [
    ("+971501234567", "+971501234567"),
    ("+971 50 123 4567", "+971501234567"),
    ("00971501234567", "+971501234567"),
])
def test_valid_phone_normalised(raw, expected):
    assert normalize_phone(raw)[1] == expected


def test_invalid_phone_row_does_not_stop_run(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [
        employee_row("EMP001", phone="0501234567", birthday=BDAY),
        employee_row("EMP002", phone="+971500000002", birthday=BDAY),
    ])
    s = run_daily(settings, TODAY, wa)
    assert len(fake_session.calls) == 1
    assert any("EMP001" in d and "Invalid phone" in d for d in s.data_errors)
    assert all("EMP002" not in d for d in s.data_errors)
    assert len(s.sent) == 1


def test_missing_phone(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [employee_row(phone=None, birthday=BDAY)])
    s = run_daily(settings, TODAY, wa)
    assert fake_session.calls == [] and any("Missing phone" in d for d in s.data_errors)


# 13. WhatsApp API failure
def test_api_failure_records_error_and_keeps_sent_year_empty(settings):
    from openpyxl import load_workbook

    settings.app_mode, settings.production_confirm = "PRODUCTION", PRODUCTION_CONFIRM_PHRASE
    make_workbook(settings.spreadsheet_path, [
        employee_row("EMP001", birthday=BDAY),
        employee_row("EMP002", phone="+971500000002", birthday=BDAY),
    ])
    sess = FakeSession([FakeResponse(400, {"error": {"code": 131026, "message": "Message undeliverable"}})])
    s = run_daily(settings, TODAY, svc(settings, sess))
    assert len(s.failed) == 1 and len(s.sent) == 1  # second employee still processed
    ws = load_workbook(settings.spreadsheet_path)["Employees"]
    header = [c.value for c in ws[1]]
    get = lambda r, c: ws.cell(row=r, column=header.index(c) + 1).value  # noqa: E731
    assert get(2, "Birthday_Sent_Year") is None
    assert get(2, "Send_Status") == "FAILED"
    assert "131026" in get(2, "Error_Message")
    assert get(3, "Birthday_Sent_Year") == 2026


def test_missing_template_error(settings, wa):
    sess = FakeSession([FakeResponse(404, {"error": {"code": 132001, "message": "Template name does not exist"}})])
    r = svc(settings, sess).send_birthday_message("+971500000001", "A")
    assert not r.success and "TEMPLATE ERROR" in r.error and not r.fatal


def test_expired_token_stops_run_and_hides_token(settings):
    make_workbook(settings.spreadsheet_path, [
        employee_row("EMP001", birthday=BDAY), employee_row("EMP002", birthday=BDAY),
    ])
    msg = f"Error validating access token: Session has expired. token={settings.whatsapp_access_token}"
    sess = FakeSession([FakeResponse(401, {"error": {"code": 190, "message": msg}})])
    s = run_daily(settings, TODAY, svc(settings, sess))
    assert s.aborted and "AUTH ERROR" in s.aborted
    assert len(sess.calls) == 1  # stopped, did not hammer the API
    assert settings.whatsapp_access_token not in s.aborted
    log_text = settings.log_file_path.read_text() if settings.log_file_path.exists() else ""
    assert settings.whatsapp_access_token not in log_text


def test_rate_limit_is_retried(settings):
    sess = FakeSession([FakeResponse(429, {"error": {"code": 130429, "message": "Rate limit hit"}})])
    r = svc(settings, sess).send_birthday_message("+971500000001", "A")
    assert r.success and len(sess.calls) == 2


def test_rate_limit_gives_up_after_max_retries(settings):
    rl = lambda: FakeResponse(429, {"error": {"code": 130429, "message": "Rate limit hit"}})  # noqa: E731
    sess = FakeSession([rl(), rl(), rl()])
    r = svc(settings, sess).send_birthday_message("+971500000001", "A")
    assert not r.success and "RATE LIMIT" in r.error and len(sess.calls) == 3


def test_spreadsheet_unavailable(settings, wa):
    s = run_daily(settings, TODAY, wa)  # no workbook created
    assert s.aborted and "Spreadsheet unavailable" in s.aborted


def test_corrupt_spreadsheet(settings, wa):
    settings.spreadsheet_path.write_text("not an excel file")
    s = run_daily(settings, TODAY, wa)
    assert s.aborted and "Spreadsheet unavailable" in s.aborted


# 14. Test mode recipient override
def test_test_mode_overrides_every_recipient(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [
        employee_row("EMP001", phone="+971500000001", birthday=BDAY),
        employee_row("EMP002", phone="+971500000002", joining=date(2016, 10, 5)),
        employee_row("EMP003", phone="+447700900123", birthday=BDAY),
    ])
    run_daily(settings, TODAY, wa)
    assert len(fake_session.calls) == 3
    assert set(fake_session.recipients) == {TEST_PHONE.lstrip("+")}


def test_test_mode_override_inside_service(settings, wa, fake_session):
    wa.send_anniversary_message("+971555555555", "CEO", 10)
    assert fake_session.recipients == [TEST_PHONE.lstrip("+")]


def test_unknown_mode_is_rejected(settings):
    settings.app_mode = "PROD"
    with pytest.raises(ConfigError):
        settings.validate()


def test_production_requires_explicit_confirmation(settings):
    settings.app_mode = "PRODUCTION"
    with pytest.raises(ConfigError, match="PRODUCTION_CONFIRM"):
        settings.validate()
    settings.production_confirm = PRODUCTION_CONFIRM_PHRASE
    settings.validate()


def test_test_mode_requires_test_phone(settings):
    settings.test_phone_number = ""
    with pytest.raises(ConfigError, match="TEST_PHONE_NUMBER"):
        settings.validate()


# 15. Dry run
def test_dry_run_sends_nothing_and_writes_nothing(settings, fake_session, caplog):
    settings.dry_run = True
    settings.whatsapp_access_token = ""  # dry run must not even need a token
    make_workbook(settings.spreadsheet_path, [
        employee_row("EMP001", name="Ahmed", birthday=BDAY),
        employee_row("EMP002", name="Sara", joining=date(2022, 10, 5)),
    ])
    before = settings.spreadsheet_path.read_bytes()
    with caplog.at_level("INFO", logger="7x"):
        s = run_daily(settings, TODAY, svc(settings, fake_session))
    assert fake_session.calls == []
    assert settings.spreadsheet_path.read_bytes() == before  # Sent Year untouched
    assert len(s.would_send) == 2
    text = caplog.text
    assert "Would send Birthday to Ahmed" in text
    assert "Would send Anniversary to Sara - 4 years" in text
    # A real run afterwards still sends (dry run did not consume the events)
    settings.dry_run = False
    settings.whatsapp_access_token = "x"
    s2 = run_daily(settings, TODAY, svc(settings, fake_session))
    assert len(s2.sent) == 2


def test_kill_switch_stop_file(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY)])
    settings.stop_file_path.write_text("stop")
    s = run_daily(settings, TODAY, wa)
    assert fake_session.calls == [] and "stopped" in s.aborted


def test_kill_switch_env(settings, wa, fake_session):
    make_workbook(settings.spreadsheet_path, [employee_row(birthday=BDAY)])
    settings.automation_enabled = False
    s = run_daily(settings, TODAY, wa)
    assert fake_session.calls == [] and "AUTOMATION_ENABLED" in s.aborted


def test_settings_repr_hides_token():
    s = Settings(whatsapp_access_token="SUPERSECRET")
    assert "SUPERSECRET" not in repr(s) and "SUPERSECRET" not in s.describe()


def test_mask_phone():
    assert mask_phone("+971501234567") == "+97150****567"


def test_arabic_template_only_when_enabled(settings, fake_session):
    w = svc(settings, fake_session)
    w.send_birthday_message("+971500000001", "Omar", language="AR")
    assert fake_session.calls[-1]["json"]["template"]["name"] == "employee_birthday"
    settings.arabic_enabled = True
    w.send_birthday_message("+971500000001", "Omar", language="AR")
    tpl = fake_session.calls[-1]["json"]["template"]
    assert tpl["name"] == "employee_birthday_ar" and tpl["language"]["code"] == "ar"


def test_log_file_never_contains_token(settings):
    import logging

    from app.main import setup_logging

    setup_logging(settings, verbose=False)
    logging.getLogger("7x").error("API said: bad token %s", settings.whatsapp_access_token)
    for h in logging.getLogger("7x").handlers:
        h.flush()
    text = settings.log_file_path.read_text()
    assert "REDACTED" in text and settings.whatsapp_access_token not in text
    logging.getLogger("7x").handlers.clear()
